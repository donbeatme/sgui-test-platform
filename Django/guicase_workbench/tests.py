import copy
import io
import json
import tempfile
from pathlib import Path
from unittest.mock import patch
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.test import TestCase, override_settings
from django.core.files.uploadedfile import SimpleUploadedFile
from rest_framework.test import APIClient
from rest_framework.exceptions import ValidationError
from projects.models import Project, ProjectMember
from testcases.models import TestCase as FunctionalCase, TestCaseModule
from langgraph_integration.models import LLMConfig
from .models import GenerationJob, CaseExtension
from .services import configuration, normalize_cases, inspect_cases, commit_job, canonical_cases
from .tasks import generate_job, split_context


def fixture(**extra):
    return {'tc_sno': 'SGUI-TC-31200000001', 'tc_name': '01 P1 登录成功', 'module_path': 'SGUI/登录',
            'precondition': '账号启用', 'priority': '中', 'steps': [{'seq': 1, 'desc': '输入用户名',
            'expected': '用户名显示在输入框', 'action': 'input', 'target': '账号输入框', 'value': 'demo',
            'locator': '#username', 'custom_step': {'kept': True}}], 'custom_case': {'kept': True}, **extra}


def ai_reply(config, prompt, content, **kwargs):
    if '需求业务分析师' in prompt:
        return json.dumps({'units': [{'id': u['id'], 'topics': [u['title']], 'summary': '业务规则与接续关系'} for u in json.loads(content)['units']]})
    if '测试架构师' in prompt:
        groups = json.loads(content)
        return json.dumps({'title': '订单出票规则验收', 'shared_group_ids': [], 'modules': [
            {'path': 'SGUI/登录', 'description': '登录业务', 'group_ids': [g['id'] for g in groups]}]})
    if '共享资料' in prompt:
        return '公共业务规则'
    body = json.loads(content.split('\n【各模块通用规则】')[0])
    refs = [m['id'] for m in body['materials']]
    return json.dumps({'cases': [fixture(source_unit_ids=refs)]})


class WorkbenchTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_superuser('wb_admin', password='test-only-password')
        self.project = Project.objects.create(name='Workbench tests', creator=self.user)
        self.other = Project.objects.create(name='Other project', creator=self.user)
        self.config = configuration(self.project)
        self.api = APIClient()
        self.api.force_authenticate(self.user)
        self.base = f'/api/projects/{self.project.pk}/workbench/'

    def new_job(self, cases=None):
        return GenerationJob.objects.create(project=self.project, creator=self.user, status='draft',
            options={'rules': self.config.rules}, cases=normalize_cases(cases or [fixture()]))

    def test_roundtrip_uses_canonical_fields_and_preserves_extensions(self):
        job = self.new_job()
        saved = commit_job(job.pk, self.project, self.user, job.revision)
        case = FunctionalCase.objects.get(pk=saved.saved_ids[0])
        self.assertEqual(case.level, 'P1')
        self.assertEqual(case.review_status, 'pending_review')
        self.assertEqual(str(case.module), 'SGUI > 登录')
        case.name = '修改后的正式名称'
        case.save()
        exported = canonical_cases(saved)[0]
        self.assertEqual(exported['tc_name'], '修改后的正式名称')
        self.assertTrue(exported['custom_case']['kept'])
        self.assertTrue(exported['steps'][0]['custom_step']['kept'])
        self.assertEqual(exported['steps'][0]['locator'], '#username')

    def test_repeated_commit_and_import_are_idempotent(self):
        job = self.new_job()
        saved = commit_job(job.pk, self.project, self.user, job.revision)
        again = commit_job(job.pk, self.project, self.user, job.revision)
        self.assertEqual(saved.saved_ids, again.saved_ids)
        changed = self.new_job([fixture(tc_name='不得覆盖正式用例')])
        commit_job(changed.pk, self.project, self.user, changed.revision)
        self.assertEqual(FunctionalCase.objects.count(), 1)
        self.assertEqual(FunctionalCase.objects.first().name, '01 P1 登录成功')

    def test_bad_expected_result_blocks_entire_batch(self):
        bad = fixture(tc_sno='bad', steps=[{'seq': 1, 'desc': '点击', 'expected': ''}])
        job = self.new_job([fixture(), bad])
        with self.assertRaises(ValidationError):
            commit_job(job.pk, self.project, self.user, job.revision)
        self.assertEqual(FunctionalCase.objects.count(), 0)
        self.assertEqual(TestCaseModule.objects.count(), 0)

    def test_disabled_rules_do_not_disable_structural_validation(self):
        cases = normalize_cases([fixture(steps=[])])
        self.assertGreater(inspect_cases(cases, [])['errors'], 0)

    def test_revision_conflict_blocks_save(self):
        job = self.new_job()
        with self.assertRaises(ValidationError):
            commit_job(job.pk, self.project, self.user, 999)
        self.assertEqual(FunctionalCase.objects.count(), 0)

    def test_deep_modules_fold_with_original_path_preserved_and_duplicate_ids_rejected(self):
        result = normalize_cases([fixture(module_path='a/b/c/d/e/f')])[0]
        self.assertEqual(result['module_path'], 'a/b/c/d/e · f')
        self.assertEqual(result['original_module_path'], 'a/b/c/d/e/f')
        with self.assertRaises(ValidationError):
            normalize_cases([fixture(), fixture()])

    def test_permissions_require_membership_and_model_permission(self):
        user = get_user_model().objects.create_user('ordinary', password='test-only-password')
        user.user_permissions.add(Permission.objects.get(codename='view_testcase', content_type__app_label='testcases'))
        ProjectMember.objects.create(project=self.other, user=user, role='member')
        self.api.force_authenticate(user)
        self.assertEqual(self.api.get(self.base).status_code, 403)
        ProjectMember.objects.create(project=self.project, user=user, role='member')
        self.assertEqual(self.api.get(self.base).status_code, 200)
        self.assertEqual(self.api.post(self.base+'preset/', {}, format='json').status_code, 403)
        self.api.force_authenticate(None)
        self.assertEqual(self.api.get(self.base).status_code, 401)

    def test_foreign_document_rejected(self):
        from requirements.models import RequirementDocument
        LLMConfig.objects.create(config_name='text', name='test', api_url='https://example.invalid/v1', is_active=True)
        doc = RequirementDocument.objects.create(project=self.other, title='Private', document_type='txt', content='secret')
        result = self.api.post(self.base+'generate/', {'text': '需求', 'document_ids': json.dumps([str(doc.pk)])})
        self.assertEqual(result.status_code, 400)
        self.assertFalse(GenerationJob.objects.exists())

    def test_text_generation_without_images_uses_server_model(self):
        LLMConfig.objects.create(config_name='text', name='test', api_url='https://example.invalid/v1', is_active=True, api_key='never-expose')
        options = {'rules': self.config.rules, 'prompts': self.config.prompts, 'stages': {}}
        job = GenerationJob.objects.create(project=self.project, creator=self.user, input={'text': '登录规则'}, options=options)
        with patch('guicase_workbench.tasks.invoke', side_effect=ai_reply) as call:
            generate_job(str(job.pk))
        job.refresh_from_db()
        self.assertEqual(job.status, 'draft')
        self.assertEqual(call.call_count, 2)
        self.assertEqual(FunctionalCase.objects.count(), 0)

    def test_multimodal_message_contains_image_and_rules_run_afterwards(self):
        LLMConfig.objects.create(config_name='vision', name='test', api_url='https://example.invalid/v1', is_active=True, supports_vision=True)
        options = {'rules': self.config.rules, 'prompts': self.config.prompts, 'stages': {}}
        job = GenerationJob.objects.create(project=self.project, creator=self.user, input={'text': '登录规则',
            'images': [{'path': 'server-owned.png', 'name': 'login.png', 'mime': 'image/png'}]}, options=options)
        with patch('guicase_workbench.tasks.default_storage.open', return_value=io.BytesIO(b'image-fixture')), \
             patch('guicase_workbench.tasks.invoke', side_effect=lambda config, prompt, content, **kwargs: '登录页有账号输入框' if isinstance(content, list) else ai_reply(config, prompt, content, **kwargs)) as call:
            generate_job(str(job.pk))
        job.refresh_from_db()
        self.assertEqual(job.status, 'draft')
        self.assertEqual(call.call_args_list[0].args[2][1]['type'], 'image_url')
        self.assertEqual(job.cases[0]['source_images'], ['login.png'])
        self.assertEqual(job.issues['errors'], 0)

    def test_vision_rejected_for_text_only_model(self):
        from .tasks import resolve_model
        LLMConfig.objects.create(config_name='text', name='test', api_url='https://example.invalid/v1', is_active=True)
        with self.assertRaises(ValidationError):
            resolve_model({}, 'vision', True)

    def test_fifty_image_uploads_accepted_and_fifty_one_rejected_before_storage(self):
        from PIL import Image
        LLMConfig.objects.create(config_name='vision', name='test',
            api_url='https://example.invalid/v1', is_active=True, supports_vision=True)
        image = io.BytesIO()
        Image.new('RGB', (2, 2)).save(image, format='PNG')
        def files(count):
            return [SimpleUploadedFile(f'page-{i}.png', image.getvalue(), content_type='image/png') for i in range(count)]
        with tempfile.TemporaryDirectory() as media, override_settings(MEDIA_ROOT=media), \
                patch('guicase_workbench.views.generate_job.delay'):
            accepted = self.api.post(self.base+'generate/', {'images': files(50)}, format='multipart')
            self.assertEqual(accepted.status_code, 201, accepted.data)
            self.assertEqual(len(GenerationJob.objects.get().input['images']), 50)
            rejected = self.api.post(self.base+'generate/', {'images': files(51)}, format='multipart')
            self.assertEqual(rejected.status_code, 400, rejected.data)
            self.assertEqual(GenerationJob.objects.count(), 1)
            self.assertEqual(len(list(Path(media).rglob('*.png'))), 50)

    def test_image_batches_resume_without_repeating_completed_images(self):
        import base64
        from .views import job_data
        LLMConfig.objects.create(config_name='vision', name='test',
            api_url='https://example.invalid/v1', is_active=True, supports_vision=True)
        assets = [{'name': f'page-{i}.png', 'path': f'page-{i}.png', 'mime': 'image/png'} for i in range(50)]
        job = GenerationJob.objects.create(project=self.project, creator=self.user,
            input={'text': '登录规则', 'images': assets},
            options={'rules': self.config.rules, 'prompts': self.config.prompts, 'stages': {}})
        image_calls = []
        fail = True
        def respond(config, prompt, content, **kwargs):
            if isinstance(content, list):
                names = [base64.b64decode(p['image_url']['url'].split(',', 1)[1]).decode()
                         for p in content if p['type'] == 'image_url']
                image_calls.append(names)
                if fail and names[0] == 'page-6.png':
                    raise RuntimeError('provider unavailable')
                return '登录页面可见账号输入框。来源：' + ', '.join(names)
            return ai_reply(config, prompt, content, **kwargs)
        with patch('guicase_workbench.pipeline.default_storage.open', side_effect=lambda path, mode: io.BytesIO(path.encode())), \
                patch('guicase_workbench.tasks.invoke', side_effect=respond):
            generate_job(str(job.pk), str(job.run_token))
            job.refresh_from_db()
            self.assertEqual(job.status, 'failed')
            self.assertEqual(len(job.progress['vision_outputs']), 1)
            first_units = copy.deepcopy([u for u in job.progress['units'] if u['kind'] == 'vision'])
            self.assertEqual(job_data(job)['progress']['vision'],
                {'image_total': 50, 'image_done': 6, 'batch_total': 9, 'batch_done': 1})
            fail = False
            image_calls.clear()
            with patch('guicase_workbench.views.generate_job.delay'):
                response = self.api.post(self.base+f'jobs/{job.pk}/resume/', {}, format='json')
                self.assertEqual(response.status_code, 200, response.data)
            job.refresh_from_db()
            generate_job(str(job.pk), str(job.run_token))
        job.refresh_from_db()
        self.assertEqual(job.status, 'draft', job.notes)
        self.assertEqual([len(c) for c in image_calls], [6, 6, 6, 6, 6, 6, 6, 2])
        self.assertEqual([n for c in image_calls for n in c], [a['name'] for a in assets[6:]])
        vision_units = [u for u in job.progress['units'] if u['kind'] == 'vision']
        self.assertEqual(vision_units[:len(first_units)], first_units)
        self.assertEqual([name for u in vision_units for name in u['source_images']], [a['name'] for a in assets])
        self.assertEqual(job_data(job)['progress']['vision']['image_done'], 50)
        self.assertNotIn('data:image', json.dumps(job.progress))
        self.assertFalse(FunctionalCase.objects.exists())

    def test_images_only_attached_to_cases_citing_visual_material(self):
        LLMConfig.objects.create(config_name='vision', name='test',
            api_url='https://example.invalid/v1', is_active=True, supports_vision=True)
        job = GenerationJob.objects.create(project=self.project, creator=self.user,
            input={'text': '登录规则', 'images': [{'path': 'login.png', 'name': 'login.png', 'mime': 'image/png'}]},
            options={'rules': self.config.rules, 'prompts': self.config.prompts, 'stages': {}})
        def respond(config, prompt, content, **kwargs):
            if isinstance(content, list):
                return '账号输入框在登录按钮上方。'
            if '测试架构师' in prompt:
                return ai_reply(config, prompt, content, **kwargs)
            materials = json.loads(content.split('\n【各模块通用规则】')[0])['materials']
            return json.dumps({'cases': [fixture(tc_sno='doc', source_unit_ids=[materials[0]['id']], source_images=['login.png']),
                fixture(tc_sno='ui', tc_name='02 P1 登录界面', source_unit_ids=[materials[1]['id']], source_images=['invented.png'])]})
        with patch('guicase_workbench.pipeline.default_storage.open', side_effect=lambda *args: io.BytesIO(b'image')), \
                patch('guicase_workbench.tasks.invoke', side_effect=respond):
            generate_job(str(job.pk))
        job.refresh_from_db()
        self.assertEqual(job.status, 'draft', job.notes)
        self.assertEqual(job.cases[0]['source_images'], [])
        self.assertEqual(job.cases[1]['source_images'], ['login.png'])

    def test_failure_does_not_fabricate_offline_success_or_expose_key(self):
        LLMConfig.objects.create(config_name='text', name='test', api_url='https://example.invalid/v1', is_active=True)
        job = GenerationJob.objects.create(project=self.project, creator=self.user, input={'text': '登录'},
            options={'rules': [], 'prompts': self.config.prompts, 'stages': {}})
        with patch('guicase_workbench.tasks.invoke', side_effect=RuntimeError('api_key=SECRET')):
            generate_job(str(job.pk))
        job.refresh_from_db()
        self.assertEqual(job.status, 'failed')
        self.assertNotIn('SECRET', json.dumps(job.notes))
        self.assertFalse(job.cases)

    def test_cancelled_job_never_invokes_model(self):
        job = self.new_job()
        job.status = 'cancelled'
        job.save()
        with patch('guicase_workbench.tasks.invoke') as call:
            generate_job(str(job.pk))
        call.assert_not_called()

    def test_preset_and_exports(self):
        payload = json.loads((Path(__file__).parent/'presets/sgui_booking_ticketing.json').read_text())
        cases = normalize_cases(payload)
        self.assertEqual(len(cases), 43)
        self.assertEqual(inspect_cases(cases, self.config.rules)['errors'], 0)
        job = self.new_job()
        from openpyxl import load_workbook
        from docx import Document
        for fmt in ['json', 'xlsx', 'docx', 'feature']:
            result = self.api.get(self.base+f'jobs/{job.pk}/export/?file_type={fmt}')
            self.assertEqual(result.status_code, 200)
            content = result.content
            if fmt == 'json':
                self.assertEqual(json.loads(content)['cases'][0]['tc_sno'], fixture()['tc_sno'])
            elif fmt == 'xlsx':
                self.assertGreater(load_workbook(io.BytesIO(content)).active.max_row, 1)
            elif fmt == 'docx':
                self.assertGreater(len(Document(io.BytesIO(content)).paragraphs), 0)
            else:
                self.assertIn('Feature:', content.decode())

    def test_settings_response_never_returns_model_secrets(self):
        LLMConfig.objects.create(config_name='text', name='test', api_url='https://example.invalid/v1', api_key='SECRET')
        self.assertNotIn('SECRET', self.api.get(self.base).content.decode())

    def test_rule_bounds_rejected_before_they_can_break_generation(self):
        rule = {'id': 'bad', 'field': 'steps', 'type': 'count', 'value': {'min': 'one', 'max': 4}, 'level': 'error', 'enabled': True, 'message': 'steps'}
        self.assertEqual(self.api.post(self.base+'settings/', {'rules': [rule]}, format='json').status_code, 400)

    def test_malformed_priority_returns_validation_error(self):
        for field in ['priority', 'level']:
            with self.assertRaises(ValidationError):
                normalize_cases([fixture(**{field: ['invalid']})])

    def model(self):
        return LLMConfig.objects.create(config_name='generation', name='test',
            api_url='https://example.invalid/v1', is_active=True)

    def test_multiple_documents_without_name_or_description_get_ai_title(self):
        from requirements.models import RequirementDocument
        self.model()
        files = [SimpleUploadedFile('rules.txt', 'R1：只允许已付款订单出票。'.encode()),
                 SimpleUploadedFile('acceptance.md', '# R2\n出票后状态为已出票。'.encode())]
        with tempfile.TemporaryDirectory() as media, override_settings(MEDIA_ROOT=media), \
                patch('guicase_workbench.views.generate_job.delay'):
            response = self.api.post(self.base+'generate/', {'docs': files, 'title': '  ', 'text': ''}, format='multipart')
            self.assertEqual(response.status_code, 201, response.data)
            job = GenerationJob.objects.get()
            self.assertTrue(job.options['auto_title'])
            self.assertEqual(len(job.input['document_ids']), 2)
            self.assertEqual(RequirementDocument.objects.filter(project=self.project).count(), 2)
            with patch('guicase_workbench.tasks.invoke', side_effect=ai_reply) as invoke:
                generate_job(str(job.pk))
            self.assertIn('R1', invoke.call_args.args[2])
            self.assertIn('R2', invoke.call_args.args[2])
            job.refresh_from_db()
            self.assertEqual(job.status, 'draft')
            self.assertEqual(job.title, '订单出票规则验收')
            self.assertFalse(FunctionalCase.objects.exists())

    def test_user_title_is_preserved(self):
        self.model()
        with patch('guicase_workbench.views.generate_job.delay'):
            response = self.api.post(self.base+'generate/', {'title': '  我的验收  ', 'text': '订单规则'})
        self.assertEqual(response.status_code, 201)
        job = GenerationJob.objects.get()
        with patch('guicase_workbench.tasks.invoke', side_effect=ai_reply):
            generate_job(str(job.pk))
        job.refresh_from_db()
        self.assertEqual(job.title, '我的验收')

    def test_long_documents_over_old_limit_are_all_processed(self):
        from requirements.models import RequirementDocument
        self.model()
        text = '文档起点\n' + ('付款后可以出票，未付款不可出票。\n' * 5000) + '文档终点'
        self.assertGreater(len(text), 80000)
        doc = RequirementDocument.objects.create(project=self.project, title='long.txt', document_type='txt', content=text)
        job = GenerationJob.objects.create(project=self.project, creator=self.user,
            input={'document_ids': [str(doc.pk)]}, options={'rules': self.config.rules, 'prompts': self.config.prompts, 'stages': {}, 'auto_title': True})
        materials = []
        def respond(config, prompt, content, **kwargs):
            if '测试架构师' not in prompt and '需求业务分析师' not in prompt:
                body = json.loads(content.split('\n【各模块通用规则】')[0])
                materials.extend(body['materials'])
            return ai_reply(config, prompt, content, **kwargs)
        with patch('guicase_workbench.tasks.invoke', side_effect=respond):
            generate_job(str(job.pk))
        job.refresh_from_db()
        self.assertEqual(job.status, 'draft', job.notes)
        self.assertTrue(any('文档起点' in m['text'] for m in materials))
        self.assertTrue(any('文档终点' in m['text'] for m in materials))
        self.assertGreater(len(materials), 1)
        self.assertEqual(job.progress['phase'], 'done')
        self.assertEqual(job.title, '订单出票规则验收')

    def test_split_preserves_every_character_and_bounds(self):
        for text in ['x' * 300000, ('一二三\n\n' * 9000) + '尾部']:
            parts = list(split_context(text, 40000))
            self.assertEqual(''.join(p[2] for p in parts), text)
            self.assertTrue(all(len(p[2]) <= 40000 for p in parts))

    def test_new_limit_fails_before_model_and_error_is_readable(self):
        from requirements.models import RequirementDocument
        self.model()
        doc = RequirementDocument.objects.create(project=self.project, title='large', document_type='txt', content='字' * 300001)
        job = GenerationJob.objects.create(project=self.project, creator=self.user, input={'document_ids': [str(doc.pk)]},
            options={'rules': [], 'prompts': self.config.prompts, 'stages': {}})
        with patch('guicase_workbench.tasks.invoke') as invoke:
            generate_job(str(job.pk))
        invoke.assert_not_called()
        job.refresh_from_db()
        self.assertEqual(job.status, 'failed')
        self.assertIn('300,000', job.notes[-1])
        self.assertNotIn('ErrorDetail', job.notes[-1])

    def test_empty_input_and_invalid_uploads_do_not_create_jobs(self):
        self.model()
        for data in [{}, {'title': '只有名称'}, {'docs': [SimpleUploadedFile('bad.exe', b'bad')]},
                     {'docs': [SimpleUploadedFile('empty.txt', b'')]},
                     {'docs': [SimpleUploadedFile(f'{i}.txt', b'rule') for i in range(21)]}]:
            response = self.api.post(self.base+'generate/', data, format='multipart')
            self.assertEqual(response.status_code, 400, response.data)
        self.assertFalse(GenerationJob.objects.exists())

    def test_cancel_during_document_preparation_stops_next_model_call(self):
        self.model()
        job = GenerationJob.objects.create(project=self.project, creator=self.user, input={'text': '规则' * 50000},
            options={'rules': [], 'prompts': self.config.prompts, 'stages': {}})
        def cancel(*args, **kwargs):
            GenerationJob.objects.filter(pk=job.pk).update(status='cancelled')
            return '整理后的规则'
        with patch('guicase_workbench.tasks.invoke', side_effect=cancel) as invoke:
            generate_job(str(job.pk))
        self.assertEqual(invoke.call_count, 1)
        job.refresh_from_db()
        self.assertEqual(job.status, 'cancelled')
        self.assertFalse(job.cases)

    def test_missing_plan_title_requests_repair(self):
        self.model()
        job = GenerationJob.objects.create(project=self.project, creator=self.user, input={'text': '规则'},
            options={'rules': [], 'prompts': self.config.prompts, 'stages': {}, 'auto_title': True})
        calls = 0
        def respond(config, prompt, content, **kwargs):
            nonlocal calls
            calls += 1
            if calls == 1:
                return json.dumps({'modules': [{'path': 'SGUI/登录', 'group_ids': ['G001']}]})
            return ai_reply(config, prompt, content, **kwargs)
        with patch('guicase_workbench.tasks.invoke', side_effect=respond):
            generate_job(str(job.pk))
        job.refresh_from_db()
        self.assertEqual(job.status, 'draft', job.notes)
        self.assertEqual(job.title, '订单出票规则验收')
        self.assertEqual(calls, 3)

    def test_timeout_retries_and_hides_provider_secrets(self):
        self.model()
        job = GenerationJob.objects.create(project=self.project, creator=self.user, input={'text': '登录规则'},
            options={'rules': [], 'prompts': self.config.prompts, 'stages': {}, 'auto_title': True})
        calls = 0
        def respond(config, prompt, content, **kwargs):
            nonlocal calls
            calls += 1
            if calls == 1:
                raise TimeoutError('provider secret=DO_NOT_LOG')
            return ai_reply(config, prompt, content, **kwargs)
        with patch('guicase_workbench.tasks.invoke', side_effect=respond), patch('guicase_workbench.pipeline.time.sleep'):
            generate_job(str(job.pk), str(job.run_token))
        job.refresh_from_db()
        self.assertEqual(job.status, 'draft', job.notes)
        self.assertEqual(calls, 3)
        self.assertNotIn('DO_NOT_LOG', json.dumps(job.notes))

    def test_resume_keeps_finished_batches_and_frozen_source(self):
        from requirements.models import RequirementDocument
        self.model()
        doc = RequirementDocument.objects.create(project=self.project, title='rules.txt', document_type='txt',
            content='# 登录\n' + '登录规则。' * 2000 + '\n# 出票\n' + '出票规则。' * 2000)
        job = GenerationJob.objects.create(project=self.project, creator=self.user, input={'document_ids': [str(doc.pk)]},
            options={'rules': [], 'prompts': self.config.prompts, 'stages': {}, 'auto_title': True})
        calls = 0
        def fail_second_batch(config, prompt, content, **kwargs):
            nonlocal calls
            if '测试架构师' not in prompt and '需求业务分析师' not in prompt:
                calls += 1
                if calls >= 2:
                    raise TimeoutError('unavailable')
            return ai_reply(config, prompt, content, **kwargs)
        old_token = str(job.run_token)
        with patch('guicase_workbench.tasks.invoke', side_effect=fail_second_batch), patch('guicase_workbench.pipeline.time.sleep'):
            generate_job(str(job.pk), old_token)
        job.refresh_from_db()
        self.assertEqual(job.status, 'failed')
        self.assertEqual(len(job.progress['outputs']), 1)
        self.assertFalse(job.cases)
        frozen = copy.deepcopy(job.progress['units'])
        completed = copy.deepcopy(job.progress['outputs'])
        doc.content = '修改后不应进入本次快照'
        doc.save()
        with patch('guicase_workbench.views.generate_job.delay') as enqueue:
            # Read-only requests must never start generation.
            self.api.get(self.base + f'jobs/{job.pk}/resume/')
            enqueue.assert_not_called()
            response = self.api.post(self.base + f'jobs/{job.pk}/resume/', {}, format='json')
            self.assertEqual(response.status_code, 200, response.data)
            self.assertEqual(self.api.post(self.base + f'jobs/{job.pk}/resume/', {}).status_code, 400)
            self.assertEqual(enqueue.call_count, 1)
        job.refresh_from_db()
        self.assertNotEqual(old_token, str(job.run_token))
        with patch('guicase_workbench.tasks.invoke', side_effect=ai_reply) as invoke:
            generate_job(str(job.pk), old_token)
            invoke.assert_not_called()
            generate_job(str(job.pk), str(job.run_token))
            self.assertEqual(invoke.call_count, 1)
        job.refresh_from_db()
        self.assertEqual(job.status, 'draft', job.notes)
        self.assertEqual(job.progress['units'], frozen)
        for key, cases in completed.items():
            self.assertEqual(job.progress['outputs'][key], cases)
        self.assertFalse(FunctionalCase.objects.exists())

    def test_cancel_and_resume_fences_late_model_response(self):
        import uuid
        self.model()
        job = GenerationJob.objects.create(project=self.project, creator=self.user, input={'text': '登录规则'},
            options={'rules': [], 'prompts': self.config.prompts, 'stages': {}})
        replacement = uuid.uuid4()
        def respond(config, prompt, content, **kwargs):
            GenerationJob.objects.filter(pk=job.pk).update(status='queued', run_token=replacement)
            return ai_reply(config, prompt, content, **kwargs)
        with patch('guicase_workbench.tasks.invoke', side_effect=respond):
            generate_job(str(job.pk), str(job.run_token))
        job.refresh_from_db()
        self.assertEqual(job.status, 'queued')
        self.assertEqual(job.run_token, replacement)
        self.assertNotIn('plan', job.progress)
        self.assertFalse(job.cases)

    def test_module_plan_must_cover_all_groups_and_use_existing_ids(self):
        from .structure import validate_plan
        groups = [{'id': 'G001'}, {'id': 'G002'}]
        plan = {'title': '登录与出票', 'modules': [{'path': '登录', 'group_ids': ['G001']}]}
        with self.assertRaises(ValidationError):
            validate_plan(plan, groups)
        plan['shared_group_ids'] = ['G999']
        with self.assertRaises(ValidationError):
            validate_plan(plan, groups)
        plan['shared_group_ids'] = ['G002']
        self.assertEqual(validate_plan(plan, groups)['modules'][0]['name'], '登录')

    def test_structural_splits_keep_complete_step_expected_pairs(self):
        from .structure import bounded_units
        rows = [{'TcStep': f'操作{i}' + '步' * 1200, 'TcExpectedResult': f'预期{i}' + '果' * 1200} for i in range(9)]
        unit = {'document': 'case.pdf', 'kind': 'case_table', 'pages': [1, 2], 'title': '跨页用例',
                'source_case_id': 'TC01', 'fields': {'TcSno': 'TC01'}, 'rows': rows, 'text': json.dumps(rows, ensure_ascii=False)}
        pieces = bounded_units([unit], limit=7000)
        self.assertGreater(len(pieces), 1)
        self.assertTrue(all(p['source_case_id'] == 'TC01' and p['pages'] == [1, 2] for p in pieces))
        for row in rows:
            self.assertTrue(any(json.dumps(row, ensure_ascii=False) in p['text'] for p in pieces))
        self.assertTrue(all(len(p['text']) <= 7000 for p in pieces))

    def test_prose_without_headings_is_read_before_planning(self):
        from .pipeline import Runner
        self.model()
        content = '开头：旅客管理。' + '背景说明。' * 4000 + '尾部：出票退款流程。'
        job = GenerationJob.objects.create(project=self.project, creator=self.user, status='running', input={'text': content},
            options={'rules': [], 'prompts': self.config.prompts, 'stages': {}})
        seen = []
        def respond(config, prompt, content, **kwargs):
            body = json.loads(content)
            seen.extend(u['text'] for u in body['units'])
            return ai_reply(config, prompt, content, **kwargs)
        runner = Runner(job, str(job.run_token), respond)
        runner.initialize()
        runner.enrich()
        while runner.state['phase'] == 'outline':
            runner.outline()
        self.assertTrue(any('开头：旅客管理' in value for value in seen))
        self.assertTrue(any('尾部：出票退款流程' in value for value in seen))
        self.assertEqual(len(runner.state['outline_outputs']), len(runner.state['units']))

    def test_other_member_cannot_resume_someone_elses_task(self):
        self.model()
        job = GenerationJob.objects.create(project=self.project, creator=self.user, status='failed',
            options={'rules': [], 'prompts': self.config.prompts, 'stages': {}})
        user = get_user_model().objects.create_user('other_member')
        user.user_permissions.add(*Permission.objects.filter(content_type__app_label='testcases', codename__in=['view_testcase', 'add_testcase']))
        ProjectMember.objects.create(project=self.project, user=user, role='member')
        self.api.force_authenticate(user)
        with patch('guicase_workbench.views.generate_job.delay') as invoke:
            result = self.api.post(self.base + f'jobs/{job.pk}/resume/', {})
        self.assertEqual(result.status_code, 403)
        invoke.assert_not_called()
